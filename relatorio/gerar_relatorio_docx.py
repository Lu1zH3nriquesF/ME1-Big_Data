"""Gera o relatório da ME em .docx (Word / Google Docs).

Uso:  python relatorio/gerar_relatorio_docx.py
Saída: relatorio/ME1_BigData_Relatorio.docx
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
DIR_GRAFICOS = RAIZ / "graficos"
DESTINO = Path(__file__).resolve().parent / "ME1_BigData_Relatorio.docx"
GITHUB = "https://github.com/Lu1zH3nriquesF/ME1-Big_Data"
GITHUB_GIT = "https://github.com/Lu1zH3nriquesF/ME1-Big_Data.git"


def _set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_hyperlink(paragraph, text: str, url: str):
    """Insere um hiperlink clicável (API do python-docx não expõe isso)."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    r_pr.append(r_fonts)
    r_pr.append(u)
    r_pr.append(color)
    r_pr.append(sz)
    new_run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def p(doc, texto: str, *, size=12, bold=False, italic=False, align="justify", space_after=8):
    paragrafo = doc.add_paragraph()
    alinhamentos = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    paragrafo.alignment = alinhamentos[align]
    paragrafo.paragraph_format.space_after = Pt(space_after)
    paragrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragrafo.paragraph_format.first_line_indent = Cm(1.25) if align == "justify" else Cm(0)
    run = paragrafo.add_run(texto)
    _set_run_font(run, size=size, bold=bold, italic=italic)
    return paragrafo


def heading(doc, texto: str, nivel: int):
    h = doc.add_heading(texto, level=nivel)
    for run in h.runs:
        _set_run_font(run, size=14 if nivel == 1 else 13, bold=True)
    return h


def tabela(doc, cabecalho: list[str], linhas: list[list[str]], titulo: str | None = None):
    if titulo:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(10)
        cap.paragraph_format.space_after = Pt(4)
        run = cap.add_run(titulo)
        _set_run_font(run, size=11, italic=True)

    tab = doc.add_table(rows=1 + len(linhas), cols=len(cabecalho))
    tab.style = "Table Grid"
    for i, nome in enumerate(cabecalho):
        cell = tab.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(nome)
        _set_run_font(run, size=10, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r, linha in enumerate(linhas, start=1):
        for c, valor in enumerate(linha):
            cell = tab.rows[r].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(valor)
            _set_run_font(run, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def figura(doc, arquivo: str, legenda: str):
    caminho = DIR_GRAFICOS / arquivo
    if not caminho.is_file():
        p(doc, f"[Figura não encontrada: {arquivo}]", italic=True, align="center")
        return
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.first_line_indent = Cm(0)
    run = pic.add_run()
    run.add_picture(str(caminho), width=Cm(15.5))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_after = Pt(12)
    run = cap.add_run(legenda)
    _set_run_font(run, size=11, italic=True)


def configurar_pagina(doc: Document):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(2.0)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Times New Roman"
    estilo.font.size = Pt(12)
    estilo._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def montar() -> Path:
    doc = Document()
    configurar_pagina(doc)

    # ------------------------------------------------------------------ capa
    for _ in range(3):
        doc.add_paragraph()
    p(doc, "[NOME DA INSTITUIÇÃO]", align="center", bold=True, size=14, space_after=4)
    p(doc, "Curso de Análise e Desenvolvimento de Sistemas", align="center", space_after=24)
    p(
        doc,
        "Análise Estatística de Focos de Calor no Brasil a partir de Dados de Sensoriamento Remoto da NASA (FIRMS/EOSDIS)",
        align="center",
        bold=True,
        size=16,
        space_after=24,
    )
    p(doc, "Medida de Eficiência — Big Data (Unidade I)", align="center", italic=True, space_after=6)
    p(doc, "Docente: Luiz Gomes da Cunha Neto", align="center", space_after=24)
    p(doc, "[Nome completo do integrante] — matrícula [XXXX]", align="center", space_after=4)
    p(doc, "[Incluir os demais integrantes do grupo]", align="center", space_after=36)
    p(doc, "[Cidade] — 2026", align="center", space_after=12)

    link_p = doc.add_paragraph()
    link_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    link_p.paragraph_format.first_line_indent = Cm(0)
    run = link_p.add_run("Código-fonte (GitHub): ")
    _set_run_font(run, size=12, bold=True)
    add_hyperlink(link_p, GITHUB, GITHUB)

    doc.add_page_break()

    # ------------------------------------------------------------------ 1
    heading(doc, "1. Introdução", 1)
    p(
        doc,
        "Queimadas e incêndios florestais no Brasil geram perda de vegetação, "
        "emissão de gases e impactos sobre saúde e economia. Observar esse "
        "fenômeno em escala nacional exige recorte espacial amplo e atualização "
        "frequente. Sensores em órbita — em especial o MODIS, a bordo dos "
        "satélites Terra e Aqua — registram anomalias térmicas mesmo em áreas "
        "de difícil acesso terrestre.",
    )
    p(
        doc,
        "Este trabalho simula um ciclo completo de análise de dados em Python: "
        "da importação do arquivo CSV à interpretação orientada à decisão. A "
        "fonte é o FIRMS (Fire Information for Resource Management System), "
        "serviço da NASA distribuído pelo EOSDIS/LANCE. A pergunta de pesquisa "
        "é: como se distribui a intensidade dos focos de calor detectados no "
        "Brasil neste recorte, e existe associação entre a intensidade do fogo "
        "(FRP) e o período de detecção (diurno/noturno)?",
    )

    # ------------------------------------------------------------------ 2
    heading(doc, "2. Material e método", 1)
    p(
        doc,
        "A análise foi implementada em Python, com pandas e numpy para "
        "manipulação e cálculo, e matplotlib/seaborn para visualização. O "
        "pipeline não faz requisição HTTP em tempo de execução: lê o CSV local "
        "originado do FIRMS. O código completo está publicado em repositório "
        "público, conforme exigência de anexar ou compartilhar o código-fonte.",
    )
    tabela(
        doc,
        ["Item", "Especificação"],
        [
            ["Linguagem", "Python 3"],
            ["Bibliotecas", "pandas, numpy, matplotlib, seaborn, tabulate"],
            ["Fonte", "NASA FIRMS / EOSDIS — produto MODIS NRT (MCD14DL)"],
            ["Portal", "https://firms.modaps.eosdis.nasa.gov/"],
            ["Meio de obtenção", "CSV local em dados/bruto/ (sem API em execução)"],
            ["Arquivo", "firms_MODIS_NRT_BRA_20260911_211449.csv"],
            ["Recorte espacial", "Brasil (BRA)"],
            ["Recorte temporal", "08/09/2026 a 11/09/2026"],
            ["Amostra", "6.321 registros (mínimo exigido: 100)"],
            ["Repositório", GITHUB],
        ],
        "Quadro 1 — Material utilizado",
    )
    p(
        doc,
        "A escolha do FIRMS, dentro do EOSDIS, é metodológica. A maior parte "
        "dos produtos EOSDIS é distribuída em NetCDF-4/HDF5 (grade raster "
        "multidimensional, sem variável categórica na mesma tabela). A ME "
        "exige variável numérica e categórica juntas. O FIRMS entrega CSV "
        "tabular com frp (escala de razão) e satellite (escala nominal).",
    )
    p(
        doc,
        "O produto utilizado é near real-time (NRT). O guia do usuário da NASA "
        "recomenda não empregar NRT em séries longas ou estudos de tendência, "
        "para os quais existe o produto padrão. Esta análise é um retrato "
        "transversal de quatro dias; portanto o NRT é adequado ao escopo. "
        "Execução: python main.py, após instalar as dependências de "
        "requirements.txt. Instruções detalhadas constam do README do repositório.",
    )

    # ------------------------------------------------------------------ 3
    heading(doc, "3. Item (a) — Descrição e importação dos dados (0,3 ponto)", 1)
    heading(doc, "3.1 Identificação da fonte", 2)
    p(
        doc,
        "Órgão: NASA — Fire Information for Resource Management System (FIRMS), "
        "distribuído pelo EOSDIS/LANCE. Portal: https://firms.modaps.eosdis.nasa.gov/. "
        "Arquivo: dados/bruto/firms_MODIS_NRT_BRA_20260911_211449.csv. "
        "Sensor: MODIS_NRT (satélites Terra e Aqua). Tema: detecção de focos de "
        "calor (queimadas e incêndios florestais), com medição da potência "
        "radiativa liberada (FRP).",
    )
    heading(doc, "3.2 Importação e pré-processamento", 2)
    p(
        doc,
        "A leitura é local (pd.read_csv em src/coleta.py). Não há chamada à API "
        "no pipeline entregue. A preparação (src/preparacao.py) reconstitui o "
        "timestamp data_hora a partir de acq_date e acq_time (HHMM), converte "
        "tipos numéricos com errors='coerce' e prevê a remoção de FRP ausente, "
        "FRP menor ou igual a zero e duplicatas exatas.",
    )
    tabela(
        doc,
        ["Etapa", "Registros"],
        [
            ["Originais", "6.321"],
            ["Timestamps inválidos", "0"],
            ["Removidos (FRP ausente)", "0"],
            ["Removidos (FRP zero)", "0"],
            ["Duplicatas removidas", "0"],
            ["Finais", "6.321 (100% retidos)"],
        ],
        "Tabela 1 — Log de limpeza",
    )
    p(
        doc,
        "Registros com FRP ausente não foram imputados com a média: imputar "
        "inventariaria intensidade de fogo que não foi medida. FRP igual a "
        "zero deslocaria média e mediana para baixo sem energia radiativa "
        "mensurável. Nesta extração nenhum registro caiu nessas regras, de "
        "modo que a representatividade da amostra original foi preservada.",
    )
    heading(doc, "3.3 Características estatísticas das variáveis", 2)
    p(
        doc,
        "Amostra final: 6.321 linhas × 15 colunas, acima do mínimo de 100 "
        "registros exigido pelo enunciado. Completude: 0% de ausentes nas "
        "colunas utilizadas. Ponto de atenção da documentação do produto: no "
        "MODIS a coluna confidence é quantitativa (0 a 100, escala de razão); "
        "no sensor VIIRS a homônima seria qualitativa ordinal "
        "(low/nominal/high). A classificação abaixo segue o MODIS.",
    )
    tabela(
        doc,
        ["Coluna", "Escala de medida"],
        [
            ["frp", "Razão (MW; zero = ausência de energia)"],
            ["brightness, bright_t31", "Razão (Kelvin)"],
            ["latitude, longitude", "Intervalar"],
            ["acq_date, acq_time, data_hora", "Intervalar"],
            ["satellite, instrument, version", "Nominal"],
            ["daynight", "Nominal dicotômica (D/N)"],
            ["confidence", "Razão (0–100)"],
        ],
        "Tabela 2 — Escalas de medida (item a)",
    )
    p(
        doc,
        "Variável numérica principal: frp (Fire Radiative Power, MW). Variável "
        "categórica: satellite (Terra ou Aqua). Auxiliar do item (d): daynight.",
    )

    # ------------------------------------------------------------------ 4
    heading(doc, "4. Item (b) — Estatística descritiva aprofundada (0,6 ponto)", 1)
    p(
        doc,
        "O FRP mede a intensidade energética do foco, e não apenas a "
        "ocorrência. Por isso foi escolhido como variável numérica principal. "
        "Os cálculos estão em src/descritiva.py.",
    )
    heading(doc, "4.1 Tendência central, dispersão e posição", 2)
    tabela(
        doc,
        ["Medida", "Valor"],
        [
            ["Média", "41,96 MW"],
            ["Mediana", "18,13 MW"],
            ["Variância (s²)", "12.467,90"],
            ["Desvio padrão (s)", "111,66 MW"],
            ["Coeficiente de variação", "266,12%"],
            ["Amplitude", "3.606,67 MW"],
            ["Mínimo", "2,52 MW"],
            ["Q1 (25%)", "9,71 MW"],
            ["Q2 (50% = mediana)", "18,13 MW"],
            ["Q3 / percentil 75", "38,56 MW"],
            ["Máximo", "3.609,19 MW"],
            ["IQR (Q3 − Q1)", "28,85 MW"],
        ],
        "Tabela 3 — Medidas descritivas do FRP (n = 6.321)",
    )
    heading(doc, "4.2 Moda da variável categórica", 2)
    tabela(
        doc,
        ["Categoria", "Frequência", "Percentual"],
        [
            ["Aqua (moda)", "4.522", "71,54%"],
            ["Terra", "1.799", "28,46%"],
        ],
        "Tabela 4 — Distribuição de satellite",
    )
    p(doc, "A distribuição é unimodal. A moda é Aqua, com 4.522 detecções.")
    heading(doc, "4.3 Justificativa: média versus mediana", 2)
    tabela(
        doc,
        ["Medida", "Valor"],
        [
            ["Assimetria (skewness)", "15,21"],
            ["Curtose", "355,63"],
            ["Diferença média − mediana", "23,83 MW"],
            ["Razão média / mediana", "2,31"],
            ["Outliers (Tukey, 1,5×IQR)", "632 (10,00%)"],
            ["Limite superior Tukey", "81,84 MW"],
            ["Maior valor", "3.609,19 MW"],
        ],
        "Tabela 5 — Forma da distribuição e outliers",
    )
    p(
        doc,
        "A média (41,96 MW) é 2,31 vezes a mediana (18,13 MW). A assimetria "
        "de 15,21 classifica a distribuição como fortemente assimétrica à "
        "direita: há uma cauda longa de valores altos. Pelo critério de Tukey, "
        "632 focos (10%) são outliers; o mais intenso atingiu 3.609,19 MW.",
    )
    p(
        doc,
        "A média incorpora todos os valores e, por isso, é arrastada pelos "
        "focos extremos. A mediana é posicional (o quinquagésimo centésimo da "
        "amostra ordenada) e praticamente ignora a magnitude da cauda. Em "
        "queimadas esse padrão é esperado: muitos focos de baixa potência e "
        "raros eventos de alta energia.",
    )
    p(
        doc,
        "O desvio padrão (111,66 MW) supera a própria média. O coeficiente de "
        "variação de 266% indica heterogeneidade extrema: não existe um “foco "
        "típico” estável em torno da média. A dispersão é da mesma ordem de "
        "grandeza — e maior — que o nível médio da variável.",
    )
    p(
        doc,
        "Para reportar a intensidade típica de um foco, adota-se a mediana "
        "(18,13 MW), por ser robusta a outliers. A média permanece útil para "
        "estimar energia total liberada, situação em que a magnitude dos "
        "extremos importa, mas não deve ser apresentada como valor típico.",
    )
    heading(doc, "4.4 Comparação entre satélites", 2)
    tabela(
        doc,
        ["Satélite", "n", "Média (MW)", "Mediana (MW)", "Desvio padrão", "Máximo"],
        [
            ["Aqua", "4.522", "44,65", "19,36", "115,27", "3.609,19"],
            ["Terra", "1.799", "35,19", "15,04", "101,75", "1.611,78"],
        ],
        "Tabela 6 — FRP por satélite",
    )
    p(
        doc,
        "O Aqua detectou mais focos e, neste recorte, apresentou média e "
        "mediana um pouco maiores. Hipótese a investigar — e não conclusão "
        "causal —: Terra e Aqua passam em horários diferentes, o que altera a "
        "chance de captar o pico diário de queima. Confirmar isso exigiria "
        "cruzar o horário de overpass com acq_time.",
    )

    # ------------------------------------------------------------------ 5
    heading(doc, "5. Item (c) — Visualização e análise da distribuição (0,4 ponto)", 1)
    p(
        doc,
        "Foram geradas quatro figuras (mínimo exigido: duas), com quartis "
        "destacados na Figura 1. Duas decisões técnicas merecem registro: "
        "(1) o eixo do FRP está em escala logarítmica, porque em escala linear "
        "os outliers comprimem a massa de dados contra a esquerda; (2) a "
        "frequência categórica usa barras, e não pizza, pois comparar "
        "comprimentos é perceptualmente mais preciso que comparar ângulos.",
    )
    figura(
        doc,
        "01_distribuicao_frp_boxplot_histograma.png",
        "Figura 1 — Distribuição do FRP: box plot e histograma, com Q1, Q2, Q3 e média destacados.",
    )
    p(
        doc,
        "A Figura 1 mostra concentração à esquerda (muitos focos de baixo FRP), "
        "cauda longa à direita e a média à direita da mediana, confirmando "
        "visualmente o diagnóstico do item (b).",
    )
    figura(
        doc,
        "02_frequencia_por_satelite.png",
        "Figura 2 — Frequência de detecções por satélite (moda = Aqua).",
    )
    p(
        doc,
        "A Figura 2 deixa explícito o desequilíbrio: Aqua concentra cerca de "
        "72% das detecções. As categorias não estão equilibradas.",
    )
    figura(
        doc,
        "03_frp_por_satelite.png",
        "Figura 3 — Box plots comparativos do FRP por satélite.",
    )
    p(
        doc,
        "Na Figura 3 as duas caixas se sobrepõem (mesma ordem de grandeza), "
        "com o Aqua um pouco deslocado para cima e máximos mais extremos. A "
        "assimetria aparece dentro de cada satélite, não apenas no agregado.",
    )
    figura(
        doc,
        "04_dispersao_geografica.png",
        "Figura 4 — Dispersão geográfica dos focos (longitude × latitude), com FRP na cor.",
    )
    p(
        doc,
        "A Figura 4 indica que os focos não se espalham de modo uniforme pelo "
        "território: há aglomerados. Isso sustenta priorizar recursos por "
        "região, e não pela média nacional. Limitação: o CSV veio de recorte "
        "por bounding box; pontos próximos à fronteira podem pertencer a "
        "países vizinhos.",
    )

    # ------------------------------------------------------------------ 6
    heading(doc, "6. Item (d) — Probabilidade simples e condicional (0,3 ponto)", 1)
    p(
        doc,
        "As probabilidades foram estimadas pela via frequentista (casos "
        "favoráveis dividido pelo total da amostra). O código está em "
        "src/probabilidade.py.",
    )
    heading(doc, "6.1 Definição formal dos eventos", 2)
    p(
        doc,
        "Probabilidade simples: A = o foco selecionado possui FRP acima do "
        "percentil 75 (valor de corte = 38,56 MW). Probabilidade condicional: "
        "A = a detecção ocorreu no período noturno (daynight = N); "
        "B = o foco possui FRP acima da média da amostra (41,96 MW).",
    )
    heading(doc, "6.2 Probabilidade simples — P(A)", 2)
    tabela(
        doc,
        ["Medida", "Valor"],
        [
            ["Evento", "FRP acima do percentil 75"],
            ["Valor de corte", "38,56 MW"],
            ["Casos favoráveis", "1.580"],
            ["Espaço amostral", "6.321"],
            ["P(A)", "0,2500 (25,00%)"],
            ["Valor teórico esperado", "0,2500"],
        ],
        "Tabela 7 — Probabilidade simples",
    )
    p(
        doc,
        "O valor empírico coincidiu com o teórico de 0,25: um quarto dos focos "
        "está estritamente acima do percentil 75. Pequenos desvios poderiam "
        "ocorrer se muitos valores caíssem exatamente no ponto de corte, pois "
        "a regra utilizada é a desigualdade estrita.",
    )
    heading(doc, "6.3 Probabilidade condicional — P(A|B)", 2)
    p(
        doc,
        "Condicionar em B reduz o espaço amostral: em vez de olhar os 6.321 "
        "focos, passa-se a olhar apenas o subconjunto em que o FRP supera a "
        "média — 1.437 registros. Desses, 240 são noturnos.",
    )
    tabela(
        doc,
        ["Medida", "Valor"],
        [
            ["A", "detecção noturna (daynight = N)"],
            ["B", "FRP acima da média (41,96 MW)"],
            ["P(A)", "0,2245 (22,45%)"],
            ["P(B)", "0,2273"],
            ["P(A ∩ B)", "0,0380"],
            ["P(A|B)", "0,1670 (16,70%)"],
            ["Espaço amostral reduzido", "1.437"],
            ["Casos favoráveis", "240"],
            ["Variação vs P(A)", "−0,0575 (−5,75 p.p.)"],
            ["Independentes?", "Não"],
        ],
        "Tabela 8 — Probabilidade condicional",
    )
    p(
        doc,
        "P(A|B) é menor que P(A): os eventos são dependentes. Saber que o foco "
        "é mais intenso que a média reduz em 5,75 pontos percentuais a chance "
        "de ele ter sido detectado à noite. Em termos operacionais, focos "
        "acima da média concentram-se mais no período diurno neste recorte; o "
        "plantão noturno permanece necessário (22% das detecções gerais são "
        "noturnas), mas o gatilho de alerta por FRP alto encontra mais eventos "
        "de dia.",
    )

    # ------------------------------------------------------------------ 7
    heading(doc, "7. Item (e) — Interpretação, insights e decisão (0,4 ponto)", 1)
    p(
        doc,
        "A intensidade não se espalha de modo homogêneo: a mediana (18 MW) "
        "descreve o foco comum, enquanto a média (42 MW) e o máximo (3.609 MW) "
        "descrevem a cauda. Focos intensos associam-se menos ao período "
        "noturno. O satélite Aqua responde por 71,5% das detecções neste arquivo.",
    )
    heading(doc, "7.1 Insight 1 — Poucos focos concentram a energia extrema", 2)
    p(
        doc,
        "Evidência: assimetria 15,21; 10% de outliers; máximo de 3.609 MW "
        "contra mediana de 18,13 MW. Interpretação: a maior parte das "
        "detecções é de baixa potência; a energia extrema está em uma minoria. "
        "Relevância: tratar os 6.321 pontos com a mesma prioridade dilui "
        "recurso. Priorizar a cauda (acima de Q3 = 38,56 MW, ou acima de "
        "81,84 MW) é mais eficiente.",
    )
    heading(doc, "7.2 Insight 2 — Intensidade acima da média é menos noturna", 2)
    p(
        doc,
        "Evidência: P(noturno) = 22,45%; P(noturno | FRP > média) = 16,70%. "
        "Interpretação: condicionar em fogo intenso reduz a probabilidade de "
        "noite em 5,75 pontos percentuais. Relevância: o monitoramento de "
        "picos de FRP deve estar reforçado no turno diurno, sem desativar a "
        "cobertura noturna.",
    )
    heading(doc, "7.3 Insight 3 — Aqua domina a amostra; o território não é uniforme", 2)
    p(
        doc,
        "Evidência: Aqua 4.522 (71,5%) versus Terra 1.799; Figura 4 com "
        "aglomerados geográficos. Interpretação: a estatística nacional mistura "
        "satélites e regiões; a mediana do Aqua (19,36 MW) supera a do Terra "
        "(15,04 MW). Relevância: alocar brigadas pela densidade local, não "
        "pela média Brasil, e declarar explicitamente o peso do Aqua nesta extração.",
    )
    heading(doc, "7.4 Decisões organizacionais propostas", 2)
    p(
        doc,
        "Organização de referência: IBAMA/Prevfogo ou secretaria estadual de "
        "meio ambiente. Cada decisão cita o número que a justifica.",
    )
    p(
        doc,
        "Primeira: instituir alerta prioritário quando o FRP superar 38,56 MW "
        "(percentil 75), concentrando ação nos 25% de focos mais energéticos "
        "(1.580 registros neste recorte). Segunda: reforçar a escala de plantão "
        "diurno para eventos acima da média (41,96 MW), com base em P(A|B) = "
        "16,7% contra 22,5%. Terceira: posicionar recursos segundo os "
        "aglomerados da Figura 4, e não segundo a média nacional (CV 266%). "
        "Quarta: manter o CSV bruto e o código no GitHub, para que a análise "
        "seja reproduzível sem depender da disponibilidade da API FIRMS.",
    )

    # ------------------------------------------------------------------ 8
    heading(doc, "8. Conclusão", 1)
    p(
        doc,
        "O ciclo importação–descritiva–gráficos–probabilidade–decisão mostra "
        "que, no Brasil entre 08 e 11 de setembro de 2026 (n = 6.321), a "
        "intensidade típica de um foco é da ordem da mediana (18 MW), enquanto "
        "a média (42 MW) é inflada por uma cauda extrema. Focos acima da média "
        "tendem a ser menos noturnos. Limitações: janela temporal curta (sem "
        "conclusão sazonal); produto NRT sujeito a reprocessamento; nuvens e "
        "horário de passagem podem ocultar focos reais; o bounding box pode "
        "incluir pontos de fronteira; associação estatística não estabelece "
        "causalidade. Próximo passo recomendado: repetir a análise com o "
        "produto padrão (MCD14ML) em janela sazonal e recorte por bioma.",
    )

    # ------------------------------------------------------------------ refs
    heading(doc, "Referências", 1)
    refs = [
        "NASA. MODIS/Aqua+Terra Thermal Anomalies/Fire locations 1km FIRMS V0061 NRT (MCD14DL). NASA Land, Atmosphere Near real-time Capability for EOS (LANCE) / Fire Information for Resource Management System (FIRMS), 2026. DOI: 10.5067/FIRMS/MODIS/MCD14DL.NRT.0061. Disponível em: https://firms.modaps.eosdis.nasa.gov/. Acesso em: 11 set. 2026.",
        "GIGLIO, L.; SCHROEDER, W.; JUSTICE, C. O. The collection 6 MODIS active fire detection algorithm and fire products. Remote Sensing of Environment, v. 178, p. 31-41, 2016. DOI: 10.1016/j.rse.2016.02.054.",
        "DAVIES, D. K.; ILAVAJHALA, S.; WONG, M. M.; JUSTICE, C. O. Fire Information for Resource Management System: Archiving and Distributing MODIS Active Fire Data. IEEE Transactions on Geoscience and Remote Sensing, v. 47, n. 1, p. 72-79, 2009. DOI: 10.1109/TGRS.2008.2002076.",
        "NASA. MODIS Collection 6/6.1 Active Fire Product User’s Guide. NASA Earthdata, 2023. Disponível em: https://www.earthdata.nasa.gov/. Acesso em: 17 set. 2026.",
        f"HENRIQUES, L. et al. ME1-Big_Data. GitHub, 2026. Disponível em: {GITHUB}. Acesso em: 17 set. 2026.",
    ]
    for ref in refs:
        paragrafo = doc.add_paragraph()
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragrafo.paragraph_format.left_indent = Cm(1.25)
        paragrafo.paragraph_format.first_line_indent = Cm(-1.25)
        paragrafo.paragraph_format.space_after = Pt(8)
        run = paragrafo.add_run(ref)
        _set_run_font(run, size=12)

    p(
        doc,
        "Incluir também as unidades de aprendizagem do catálogo Sagah listadas "
        "no enunciado da disciplina, conforme orientação do docente.",
        italic=True,
    )

    # ------------------------------------------------------------------ apêndice
    heading(doc, "Apêndice — Código-fonte", 1)
    p(
        doc,
        "O enunciado exige que o código seja anexado ou tenha o link "
        "compartilhado no relatório. Adota-se o repositório público abaixo, "
        "acessível sem login.",
    )
    bloco = doc.add_paragraph()
    bloco.alignment = WD_ALIGN_PARAGRAPH.LEFT
    bloco.paragraph_format.first_line_indent = Cm(0)
    run = bloco.add_run("Repositório: ")
    _set_run_font(run, size=12, bold=True)
    add_hyperlink(bloco, GITHUB, GITHUB)

    bloco2 = doc.add_paragraph()
    bloco2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    bloco2.paragraph_format.first_line_indent = Cm(0)
    run = bloco2.add_run("Clone: ")
    _set_run_font(run, size=12, bold=True)
    add_hyperlink(bloco2, GITHUB_GIT, GITHUB_GIT)

    p(doc, "Árvore principal do projeto:", align="left")
    arvore = (
        "config.py\n"
        "main.py\n"
        "validar_pipeline.py\n"
        "requirements.txt\n"
        "src/\n"
        "    coleta.py\n"
        "    preparacao.py\n"
        "    descritiva.py\n"
        "    visualizacao.py\n"
        "    probabilidade.py\n"
        "    interpretacao.py\n"
        "dados/bruto/firms_MODIS_NRT_BRA_20260911_211449.csv\n"
        "graficos/*.png\n"
        "relatorio/"
    )
    code = doc.add_paragraph()
    code.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.left_indent = Cm(1)
    run = code.add_run(arvore)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(10)

    p(
        doc,
        "Reprodução local: python -m pip install -r requirements.txt  e, em seguida,  python main.py. "
        "O CSV em dados/bruto/ precisa acompanhar o código; sem ele o pipeline não executa.",
    )

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DESTINO)
    return DESTINO


if __name__ == "__main__":
    caminho = montar()
    print(f"Arquivo gerado: {caminho}")
    print(f"Tamanho: {caminho.stat().st_size / 1024:.1f} KB")
